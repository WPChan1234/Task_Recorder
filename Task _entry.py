import datetime
import time
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import tkinter as tk
from tkinter import filedialog

root = tk.Tk()
FilePath=filedialog.askopenfilename(title="Select excel file for task recording")
root.destroy()

wb = load_workbook(filename=FilePath)
ws=wb["Sheet1"]
Table=ws.tables["Task_Record"]

for row in Table:
    ws.apprend(row)
    print(row)




